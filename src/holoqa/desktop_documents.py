"""Text-only document intake for the native desktop picker. Never renders files."""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
import signal
import stat
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from xml.etree import ElementTree

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 100_000
EXTENSIONS = {'.txt', '.md', '.markdown', '.pdf', '.docx', '.xlsx', '.xlsm'}
MAX_OUTPUT = MAX_TEXT * 6 + 8192
EXTRACTION_TIMEOUT = 12.0


class DocumentError(ValueError):
    """Actionable extraction error safe to show without local paths/content."""

    def __init__(self, message: str, code: str = 'invalid_document'):
        super().__init__(message)
        self.code = code


class DocumentCancelled(DocumentError):
    def __init__(self):
        super().__init__('Document extraction cancelled. Attach the document again when ready.', 'cancelled')


def _worker_command(path: str) -> list[str]:
    if getattr(sys, 'frozen', False):
        return [sys.executable, '--extract-document', path]
    return [sys.executable, '-B', '-m', 'holoqa.desktop_documents', '--extract', path]


def read_document(path: str, *, cancel: threading.Event | None = None,
                  timeout: float = EXTRACTION_TIMEOUT) -> dict:
    """Extract in an owned worker, bounding runtime and protocol output.

    A bounded pump spools stdout to disk rather than accumulating parser output
    in memory. Even a noisy/malfunctioning worker cannot grow the spool without
    limit. No document parser executes in the long-lived desktop engine.
    """
    if cancel is not None and cancel.is_set():
        raise DocumentCancelled()
    overflow, pump_done = threading.Event(), threading.Event()
    process = None
    pump = None
    with tempfile.TemporaryFile(mode='w+b') as spool:
        try:
            process = subprocess.Popen(_worker_command(str(path)), stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, start_new_session=os.name != 'nt',
                creationflags=0x08000000 if os.name == 'nt' else 0)

            def collect():
                total = 0
                try:
                    while chunk := process.stdout.read(8192):
                        remaining = MAX_OUTPUT + 1 - total
                        spool.write(chunk[:remaining])
                        total += len(chunk)
                        if total > MAX_OUTPUT:
                            overflow.set()
                            return
                except (OSError, ValueError):
                    overflow.set()
                finally:
                    pump_done.set()

            pump = threading.Thread(target=collect, daemon=True, name='document-output')
            pump.start()
            deadline = time.monotonic() + timeout
            while True:
                if cancel is not None and cancel.is_set():
                    raise DocumentCancelled()
                if overflow.is_set():
                    raise DocumentError('Document extraction exceeded its output limit. Attach a smaller text document.', 'output_limit')
                if time.monotonic() >= deadline:
                    raise DocumentError('Document extraction took too long. Attach a smaller or simpler text document.', 'timeout')
                if process.poll() is not None and pump_done.is_set():
                    break
                pump_done.wait(.02) if process.poll() is not None else time.sleep(.02)
            if process.returncode != 0:
                raise DocumentError('Document extraction stopped unexpectedly. Attach a smaller or simpler text document.', 'worker_failed')
            spool.seek(0)
            try:
                response = json.loads(spool.read(MAX_OUTPUT + 1).decode('utf-8'))
            except (ValueError, UnicodeError) as error:
                raise DocumentError('Document extraction returned an invalid response. Try a plain UTF-8 text document.', 'worker_protocol') from error
            if not isinstance(response, dict):
                raise DocumentError('Document extraction returned an invalid response.', 'worker_protocol')
            if 'error' in response:
                error = response['error']
                if not isinstance(error, dict):
                    raise DocumentError('Document extraction returned an invalid response.', 'worker_protocol')
                # Worker messages are our fixed strings, never parser exceptions.
                raise DocumentError(str(error.get('message', 'Document extraction failed.'))[:500],
                                    str(error.get('code', 'invalid_document'))[:50])
            result = response.get('result')
            if (not isinstance(result, dict) or not isinstance(result.get('text'), str)
                    or len(result['text']) > MAX_TEXT):
                raise DocumentError('Document extraction returned an invalid response.', 'worker_protocol')
            return result
        except OSError as error:
            raise DocumentError('Cannot start document extraction. Restart HoloQA and try again.', 'worker_start') from error
        finally:
            if process is not None:
                if os.name != 'nt':
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
                elif process.poll() is None:
                    process.kill()
                process.wait(timeout=5)
                if pump is not None:
                    pump.join(timeout=2)
                if process.stdout:
                    process.stdout.close()


def _worker_limits():
    if sys.platform.startswith('linux'):
        import resource
        # Parser CPU and virtual-memory limits supplement the parent's deadline.
        # Windows workers are bounded by time/output and owned by the shell Job.
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (8, 9))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT, MAX_OUTPUT))


def extract_main(path: str) -> int:
    """Worker entrypoint; frozen desktop dispatches --extract-document here."""
    try:
        # Resource-limited parsers must not create truncated shared .pyc files.
        sys.dont_write_bytecode = True
        _worker_limits()
        response = {'result': _read_document(path)}
    except DocumentError as error:
        response = {'error': {'code': error.code, 'message': str(error)}}
    except MemoryError:
        response = {'error': {'code': 'resource_limit', 'message': 'Document extraction exceeded its memory limit. Attach a simpler text document.'}}
    except (FileNotFoundError, PermissionError, OSError):
        response = {'error': {'code': 'unavailable', 'message': 'Document cannot be read. Check access and select the attachment again.'}}
    except Exception:
        # Parsers include filenames and document fragments in exceptions. Never
        # forward those diagnostics onto the desktop protocol or model prompt.
        response = {'error': {'code': 'malformed', 'message': 'Document is malformed or unsupported. Export it as UTF-8 text and attach again.'}}
    payload = json.dumps(response, ensure_ascii=False, separators=(',', ':'))
    if len(payload.encode('utf-8')) > MAX_OUTPUT:
        payload = '{"error":{"code":"output_limit","message":"Extracted document output is too large."}}'
    sys.stdout.buffer.write(payload.encode('utf-8'))
    sys.stdout.buffer.flush()
    return 0


def _read_document(path: str) -> dict:
    candidate = Path(path)
    if candidate.suffix.lower() not in EXTENSIONS:
        raise DocumentError('Supported documents: TXT, Markdown, PDF, DOCX, and XLSX.', 'unsupported_type')
    if not stat.S_ISREG(candidate.stat().st_mode):
        raise DocumentError('Attachments must be regular files.', 'invalid_type')
    with candidate.open('rb') as handle:
        data = handle.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise DocumentError('Each document must be 5 MiB or smaller.', 'source_limit')
    suffix = candidate.suffix.lower()
    if suffix in {'.xlsx', '.xlsm'}:
        # Read values only. Macros, formulas, links, images, and embedded
        # objects are intentionally ignored. Preflight the ZIP before the
        # workbook parser sees it to bound expansion from hostile archives.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(max(0, e.file_size) for e in entries) > 32 * 1024 * 1024:
                raise DocumentError('Excel workbook exceeds the extraction limits.', 'archive_limit')
            if any(e.file_size > 8 * 1024 * 1024 for e in entries):
                raise DocumentError('Excel workbook contains an oversized sheet.', 'archive_limit')
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True,
                                    keep_links=False, keep_vba=False)
            rows, cells, chunks = 0, 0, []
            for sheet in workbook.worksheets:
                chunks.append(f'## Sheet: {sheet.title}')
                for row in sheet.iter_rows(values_only=True):
                    rows += 1
                    if rows > 20_000:
                        raise DocumentError('Excel workbooks are limited to 20,000 rows.', 'row_limit')
                    values = []
                    for value in row:
                        cells += 1
                        if cells > 200_000:
                            raise DocumentError('Excel workbooks are limited to 200,000 cells.', 'cell_limit')
                        values.append('' if value is None else str(value).replace('\x00', ''))
                    line = ' | '.join(values).rstrip()
                    if line:
                        chunks.append(line)
            workbook.close()
            text = '\n'.join(chunks)
        except DocumentError:
            raise
        except Exception as error:
            raise DocumentError('Excel workbook is malformed or unsupported. Export it as .xlsx and attach again.', 'malformed') from error
    elif suffix == '.pdf':
        from pypdf import PdfReader, overwrite_configuration
        # Bound individual compressed streams as well as file/page/text sizes.
        overwrite_configuration(
            maximum_declared_stream_length=8 * 1024 * 1024,
            array_based_stream_maximum_output_length=8 * 1024 * 1024,
            zlib_maximum_output_length=8 * 1024 * 1024,
            lzw_maximum_output_length=8 * 1024 * 1024,
            run_length_maximum_output_length=8 * 1024 * 1024,
            brotli_maximum_output_length=8 * 1024 * 1024,
            stream_decoding_work_maximum_length=16 * 1024 * 1024,
            stream_parsing_work_maximum_count=100_000,
        )
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentError('Encrypted PDFs are not supported. Attach an unlocked text document.', 'encrypted')
        if len(reader.pages) > 100:
            raise DocumentError('PDF attachments are limited to 100 pages.', 'page_limit')
        chunks, length = [], 0
        for page in reader.pages:
            text = page.extract_text() or ''
            length += len(text) + 1
            if length > MAX_TEXT:
                raise DocumentError('Extracted document text exceeds 100,000 characters.', 'text_limit')
            chunks.append(text)
        text = '\n'.join(chunks)
    elif suffix == '.docx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 16 * 1024 * 1024:
                raise DocumentError('DOCX archive exceeds the extraction limits.', 'archive_limit')
            entry = archive.getinfo('word/document.xml')
            if entry.file_size > 2 * 1024 * 1024:
                raise DocumentError('DOCX document XML exceeds 2 MiB.', 'archive_limit')
            xml = archive.read(entry)
            tree = ElementTree.fromstring(xml)
            ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            text = '\n'.join(''.join(p.itertext()) for p in tree.iter(ns + 'p'))
    else:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError as error:
            raise DocumentError('Text attachments must use UTF-8 encoding.', 'encoding') from error
    if not text.strip():
        raise DocumentError('No readable text found in this attachment.', 'no_text')
    if '\x00' in text:
        raise DocumentError('Binary content is not a text document.', 'binary_content')
    if len(text) > MAX_TEXT:
        raise DocumentError('Extracted document text exceeds 100,000 characters.', 'text_limit')
    return {'name': candidate.name, 'size': len(data), 'characters': len(text),
            'sha256': hashlib.sha256(data).hexdigest(), 'text': text}


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--extract', required=True)
    raise SystemExit(extract_main(parser.parse_args().extract))
