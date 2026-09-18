"""Optional XLSX annotation.

A workbook is a deliverable, not a source of truth. HoloQA never overwrites the
original: it copies it, appends Status / Actual Result / Evidence columns, and
writes the copy beside the run.

Rows are matched by a step-id column, so this works for any checklist workbook
without per-application cell mapping. That is the one thing the prior tool could
not do — it addressed fixed cells (``C3``-``C8``, rows 12-46) and so was welded
to a single spreadsheet.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

#: Header labels accepted as the step-identifier column, case-insensitive.
ID_HEADERS = ("step id", "test id", "id", "langkah", "no")
ADDED = ("Status", "Actual Result", "Evidence")


class WorkbookError(ValueError):
    pass


def _headers(sheet) -> list[str]:
    first = next(sheet.iter_rows(values_only=True), None)
    if first is None:
        raise WorkbookError("workbook is empty")
    return [str(value or "").strip().lower() for value in first]


def annotate(source: Path, output: Path, steps: list[dict[str, Any]]) -> Path:
    """Copy ``source`` and fill in results keyed by step id."""
    source, output = Path(source), Path(output)
    if not source.is_file():
        raise WorkbookError(f"source workbook not found: {source}")

    book = load_workbook(source)
    sheet = book.active
    headers = _headers(sheet)

    id_column = next(
        (headers.index(name) for name in ID_HEADERS if name in headers), None
    )
    if id_column is None:
        raise WorkbookError(
            "no step-id column found; expected one of: " + ", ".join(ID_HEADERS)
        )

    for label in ADDED:
        if label.lower() not in headers:
            sheet.cell(row=1, column=sheet.max_column + 1, value=label)
            headers.append(label.lower())

    by_id = {str(step["id"]): step for step in steps}
    written = 0
    for row in range(2, sheet.max_row + 1):
        key = str(sheet.cell(row=row, column=id_column + 1).value or "").strip()
        step = by_id.get(key)
        if not step:
            continue
        evidence = "; ".join(item["file"] for item in step.get("observations", []))
        sheet.cell(row=row, column=headers.index("status") + 1, value=step.get("verdict") or "")
        sheet.cell(row=row, column=headers.index("actual result") + 1, value=step.get("note") or "")
        sheet.cell(row=row, column=headers.index("evidence") + 1, value=evidence)
        written += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    if written == 0:
        raise WorkbookError(
            "no workbook row matched a step id; check that the id column values "
            "match the step ids in the plan"
        )
    return output
