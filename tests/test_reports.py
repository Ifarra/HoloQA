from pathlib import Path

from openpyxl import Workbook, load_workbook

from holoqa.mcp_server import holoqa_export_report


def test_export_report_creates_html_json_and_xlsx(tmp_path: Path):
    state = tmp_path / "state.db"
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Test ID", "Title", "Steps", "Expected Result"])
    sheet.append(["MVP", "Demo", "Open", "Works"])
    workbook.save(source)

    from holoqa.runs import RunStore

    plan = RunStore(state).create_plan("project", [{"test_id": "MVP"}])
    RunStore(state).approve(plan.plan_id)
    run = RunStore(state).execute(plan.plan_id)
    result = holoqa_export_report(run.run_id, str(state), str(source))

    assert set(result["artifacts"]) == {"json", "html", "xlsx"}
    assert Path(result["artifacts"]["json"]).is_file()
    assert Path(result["artifacts"]["html"]).is_file()
    workbook = load_workbook(result["artifacts"]["xlsx"])
    assert workbook.active.cell(row=1, column=5).value == "Status"
