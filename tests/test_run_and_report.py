"""Guardrails, packaging, workbook annotation, and cross-run history."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from holoqa import BLOCKED, FAIL, PASS
from holoqa import history as history_module
from holoqa import report as report_module
from holoqa import verdict as verdict_module
from holoqa.run import GuardrailError, Run, decision
from holoqa.workbook import WorkbookError, annotate
from tests.conftest import api_evidence, dom_evidence, screenshot_evidence


def _pass_t1(plan, run):
    api_evidence(
        run, "T1", "t1-create",
        method="POST", status=201, body={"id": "thing_42", "status": "PENDING"},
    )
    screenshot_evidence(run, "T1", "t1-page")
    outcome, results, _ = verdict_module.evaluate(plan, run, "T1")
    run.set_verdict("T1", outcome, assertions=results)
    return outcome


# --------------------------------------------------------------- guardrails


def test_agent_may_assert_blocked_but_nothing_else(run):
    run.set_verdict("T1", BLOCKED, note="feature flag is off in this environment", by="agent")
    assert run.read()["steps"]["T1"]["decided_by"] == "agent"

    with pytest.raises(GuardrailError, match="may only assert"):
        run.set_verdict("T1", PASS, by="agent")


def test_evidence_must_exist_and_be_non_empty(run, tmp_path):
    with pytest.raises(GuardrailError, match="does not exist"):
        run.attach("T1", kind="screenshot", path=tmp_path / "ghost.png")

    empty = run.evidence_dir / "empty.png"
    empty.write_bytes(b"")
    with pytest.raises(GuardrailError, match="empty"):
        run.attach("T1", kind="screenshot", path=empty)


def test_dependencies_gate_judgement(plan, run):
    assert run.blocked_dependencies(plan, "T2") == ["T1"]
    _pass_t1(plan, run)
    assert run.blocked_dependencies(plan, "T2") == []


def test_revision_trail_preserves_the_superseded_note(run):
    run.set_verdict("T1", BLOCKED, note="staging was unreachable")
    run.set_verdict("T1", FAIL, note="reachable now, returns 500")
    revisions = run.read()["steps"]["T1"]["revisions"]
    assert len(revisions) == 1
    assert revisions[0]["from"] == BLOCKED
    assert revisions[0]["previous_note"] == "staging was unreachable"


def test_one_failure_holds_the_release():
    assert decision({PASS: 5, FAIL: 0, BLOCKED: 0}) == "RELEASE"
    assert decision({PASS: 5, FAIL: 1, BLOCKED: 0}) == "HOLD"
    assert decision({PASS: 5, FAIL: 0, BLOCKED: 1}) == "HOLD"


# ---------------------------------------------------------------- packaging


def test_strict_packaging_refuses_an_unjudged_run(plan, run):
    _pass_t1(plan, run)
    with pytest.raises(report_module.PackageError, match="not yet judged"):
        report_module.package(run, strict=True)


def test_package_produces_a_zip_with_report_and_evidence(plan, run):
    _pass_t1(plan, run)
    run.set_verdict("T2", BLOCKED, note="depends on an endpoint absent in staging")
    run.set_verdict("T3", BLOCKED, note="progress window closed before the second capture")

    result = report_module.package(run, strict=True)

    assert result["decision"] == "HOLD"
    assert result["counts"] == {PASS: 1, FAIL: 0, BLOCKED: 2}

    with zipfile.ZipFile(result["archive"]) as bundle:
        names = bundle.namelist()
        assert "report.md" in names
        assert "run.json" in names
        assert any(name.startswith("evidence/") for name in names)
        report = bundle.read("report.md").decode("utf-8")

    assert "decision: HOLD" in report
    assert "One failed or blocked step holds the release." in report
    assert "holoqa" in report  # provenance column


def test_report_lists_revised_verdicts(plan, run):
    _pass_t1(plan, run)
    run.set_verdict("T1", BLOCKED, note="retracted after finding a stale fixture")
    run.set_verdict("T2", BLOCKED, note="blocked by T1")
    run.set_verdict("T3", BLOCKED, note="blocked by T1")
    report_module.package(run, strict=True)
    report = (run.out_dir / "report.md").read_text(encoding="utf-8")
    assert "## Revised verdicts" in report
    assert "PASS → BLOCKED" in report


# ----------------------------------------------------------------- workbook


def test_workbook_annotation_matches_rows_by_step_id(tmp_path, plan, run):
    source = tmp_path / "checklist.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Step ID", "Title", "Expected"])
    sheet.append(["T1", "Create a thing", "201 PENDING"])
    sheet.append(["T2", "Read it back", "200 Ready"])
    book.save(source)

    _pass_t1(plan, run)
    run.set_verdict("T2", BLOCKED, note="endpoint absent in staging")
    steps = list(run.read()["steps"].values())

    output = annotate(source, tmp_path / "out.xlsx", steps)

    from openpyxl import load_workbook

    filled = load_workbook(output).active
    headers = [str(c.value or "").lower() for c in filled[1]]
    status_col = headers.index("status") + 1
    assert filled.cell(row=2, column=status_col).value == PASS
    assert filled.cell(row=3, column=status_col).value == BLOCKED
    # the original is never touched
    assert load_workbook(source).active.max_column == 3


def test_workbook_without_an_id_column_is_refused(tmp_path):
    source = tmp_path / "bad.xlsx"
    book = Workbook()
    book.active.append(["Description", "Notes"])
    book.save(source)
    with pytest.raises(WorkbookError, match="no step-id column"):
        annotate(source, tmp_path / "out.xlsx", [])


# ------------------------------------------------------------------ history


def test_compare_reports_regressions_against_the_last_green_run(tmp_path, plan):
    runs_root = tmp_path / ".holoqa" / "runs"

    green = Run.create(runs_root / "20260101-0900-fixture", plan)
    for step_id in ("T1", "T2", "T3"):
        screenshot_evidence(green, step_id, f"{step_id}-shot")
        green.set_verdict(step_id, PASS)
    history_module.record(green)

    later = Run.create(runs_root / "20260102-0900-fixture", plan)
    for step_id in ("T1", "T2", "T3"):
        screenshot_evidence(later, step_id, f"{step_id}-shot")
    later.set_verdict("T1", PASS)
    later.set_verdict("T2", FAIL, note="500 from /things")
    later.set_verdict("T3", PASS)
    history_module.record(later)

    result = history_module.compare(later)

    assert result["status"] == "ok"
    assert result["baseline"] == "20260101-0900-fixture"
    assert [item["step"] for item in result["regressions"]] == ["T2"]
    assert result["fixes"] == []


def test_compare_without_history_says_so(plan, run):
    assert history_module.compare(run)["status"] == "no_baseline"


def test_flaky_steps_surface_after_enough_runs(tmp_path, plan):
    runs_root = tmp_path / ".holoqa" / "runs"
    last = None
    for index, t2 in enumerate([PASS, FAIL, PASS, FAIL]):
        current = Run.create(runs_root / f"2026010{index}-0900-fixture", plan)
        screenshot_evidence(current, "T1", "t1-shot")
        screenshot_evidence(current, "T2", "t2-shot")
        current.set_verdict("T1", PASS)
        current.set_verdict("T2", t2, note="flaps" if t2 != PASS else "")
        history_module.record(current)
        last = current

    unstable = history_module.flaky(last)
    assert [item["step"] for item in unstable] == ["T2"]
    assert unstable[0]["runs"] == 4
