from __future__ import annotations

import base64
import io
import json
import os
import shutil
import struct
import subprocess
import urllib.request
import zipfile
import zlib
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


def _png_for_case(index: int) -> bytes:
    """Generate a visible deterministic browser-like frame without Pillow."""
    width, height = 320, 180
    pixels = bytearray()
    for y in range(height):
        pixels.append(0)
        for x in range(width):
            accent = (x // 32) == ((index - 1) % 10) and 28 < y < 152
            pixels.extend((238, 105, 54) if accent else (244, 244, 239))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(bytes(pixels), 9)) + chunk(b"IEND", b"")


def _upload(url: str, content: bytes, content_type: str) -> None:
    request = urllib.request.Request(url, data=content, method="PUT", headers={"Content-Type": content_type})
    with urllib.request.urlopen(request, timeout=20) as response:
        assert response.status in {200, 204}


def _make_fixture(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    (root / "index.html").write_text(
        "<!doctype html><html><body><main><h1>HoloQA Fixture</h1><button id='continue'>Continue</button><p id='state'>Ready</p></main>"
        "<script>document.querySelector('#continue').onclick=()=>document.querySelector('#state').textContent='Complete'</script></body></html>",
        encoding="utf-8",
    )
    (root / "README.md").write_text("Deterministic HoloQA remote-agent fixture.\n", encoding="utf-8")
    workbook_path = root / "SITUAT.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SITUAT"
    sheet.append(["Test ID", "Title", "Steps", "Expected Result"])
    for index in range(1, 11):
        sheet.append([f"SIT-{index:03d}", f"Fixture task {index}", "Open fixture; Click Continue; Verify Complete", "Fixture state is Complete"])
    workbook.save(workbook_path)

    git = ["git", "-C", str(root)]
    if shutil.which("git") and not (root / ".git").is_dir():
        subprocess.run([*git, "init"], check=True, capture_output=True)
        subprocess.run([*git, "config", "user.email", "holoqa@example.test"], check=True)
        subprocess.run([*git, "config", "user.name", "HoloQA E2E"], check=True)
    assert (root / ".git").is_dir(), "fixture must be created as a real Git repository on the host"
    if shutil.which("git"):
        subprocess.run([*git, "add", "index.html", "README.md", "SITUAT.xlsx"], check=True)
        subprocess.run([*git, "commit", "-m", "fixture"], check=True, capture_output=True)
    assert len(list(load_workbook(workbook_path, read_only=True).active.iter_rows(min_row=2, values_only=True))) == 10
    return workbook_path


def test_remote_agent_workflow_with_live_monitoring_and_garage(tmp_path: Path, monkeypatch):
    # This path is mounted into the service as /workspace in docker-compose.dev.yml.
    fixture_root = Path(os.environ.get("HOLOQA_E2E_FIXTURE", "/workspace/.holoqa-e2e-fixture"))
    workbook_path = _make_fixture(fixture_root)
    database = Path(os.environ.get("HOLOQA_E2E_STATE", str(tmp_path / "state.db")))
    monkeypatch.setenv("HOLOQA_STATE", str(database))
    monkeypatch.setenv("HOLOQA_OBJECT_STORAGE", "s3")
    s3_endpoint = os.environ.get("HOLOQA_E2E_S3_ENDPOINT", "http://garage:3900")
    public_endpoint = os.environ.get("HOLOQA_E2E_S3_PUBLIC_ENDPOINT", s3_endpoint)
    monkeypatch.setenv("HOLOQA_S3_ENDPOINT", s3_endpoint)
    monkeypatch.setenv("HOLOQA_S3_PUBLIC_ENDPOINT", public_endpoint)
    monkeypatch.setenv("HOLOQA_S3_ACCESS_KEY_ID", os.environ.get("HOLOQA_S3_ACCESS_KEY_ID", "GKholoqa"))
    monkeypatch.setenv("HOLOQA_S3_SECRET_ACCESS_KEY", os.environ.get("HOLOQA_S3_SECRET_ACCESS_KEY", "change-me-in-production"))

    from holoqa.dashboard import create_app
    from holoqa.mcp_server import (
        holoqa_approve_run,
        holoqa_complete_run,
        holoqa_create_run_plan,
        holoqa_get_run_status,
        holoqa_import_test_workbook,
        holoqa_initialize_project,
        holoqa_open_run_session,
        holoqa_prepare_artifact_upload,
        holoqa_record_agent_event,
        holoqa_update_case_verdict,
    )
    from holoqa.object_storage import artifact_exists
    from holoqa.project_store import ProjectStore
    from holoqa.runs import RunStore

    initialized = holoqa_initialize_project(str(fixture_root))
    project_id = initialized["project_id"]
    imported = holoqa_import_test_workbook(str(workbook_path), project_id)
    assert imported["case_count"] == 10
    assert len(ProjectStore(database).testcases(project_id)) == 10

    plan = holoqa_create_run_plan(project_id)
    assert len(plan["cases"]) == 10
    holoqa_approve_run(plan["plan_id"])
    session = holoqa_open_run_session(plan["plan_id"], "e2e-ai-coder")
    run_id = session["run_id"]
    token = session["session_token"]

    app = create_app(database)
    with TestClient(app) as client:
        with client.websocket_connect(f"/api/runs/{run_id}/live?role=viewer") as viewer:
            initial = viewer.receive_json()
            assert initial["type"] == "live_state"
            assert initial["connected"] is False
            with client.websocket_connect(f"/api/runs/{run_id}/live?role=agent&token={token}") as agent:
                agent.send_json({"type": "heartbeat"})
                assert viewer.receive_json()["type"] == "agent_state"
                agent.send_json({"type": "frame", "data": base64.b64encode(_png_for_case(1)).decode(), "mime": "image/png", "url": "http://fixture.local/"})
                frame = viewer.receive_json()
                assert frame["type"] == "frame"
                assert frame["mime"] == "image/png"
                agent.send_json({"type": "event", "sequence": 1, "event_type": "agent_observation", "test_id": "SIT-001", "message": "Fixture page observed"})
                live_event = viewer.receive_json()
                assert live_event["type"] == "agent_event"
                assert live_event["event"]["sequence"] == 1
            disconnected = viewer.receive_json()
            assert disconnected["type"] == "agent_state"
            assert disconnected["connected"] is False

    sequence = 2
    evidence_uris: list[str] = []
    for index in range(1, 11):
        test_id = f"SIT-{index:03d}"
        event = holoqa_record_agent_event(run_id, token, sequence, "case_started", test_id=test_id, message=f"Starting {test_id}")
        assert event["last_event_sequence"] == sequence
        sequence += 1
        screenshot = holoqa_prepare_artifact_upload(run_id, token, test_id, f"{test_id}.png", "image/png")
        _upload(screenshot["upload_url"], _png_for_case(index), "image/png")
        artifacts = [{"artifact_uri": screenshot["artifact_uri"], "kind": "screenshot"}]
        supporting = [
            (f"{test_id}.dom.html", "text/html", f"<main data-test-id=\"{test_id}\"><h1>HoloQA Fixture</h1><p>Complete</p></main>".encode(), "dom"),
            (f"{test_id}.console.log", "text/plain", f"[{test_id}] no console errors\n".encode(), "console"),
            (f"{test_id}.har", "application/json", json.dumps({"log": {"version": "1.2", "entries": [{"request": {"url": "http://fixture.local/"}}]}}).encode(), "network"),
        ]
        trace = io.BytesIO()
        with zipfile.ZipFile(trace, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("trace.txt", f"Trace for {test_id}: open -> click -> verify\n")
        supporting.append((f"{test_id}.trace.zip", "application/zip", trace.getvalue(), "trace"))
        for filename, content_type, content, kind in supporting:
            artifact = holoqa_prepare_artifact_upload(run_id, token, test_id, filename, content_type)
            _upload(artifact["upload_url"], content, content_type)
            artifacts.append({"artifact_uri": artifact["artifact_uri"], "kind": kind})
            evidence_uris.append(artifact["artifact_uri"])
        evidence_uris.append(screenshot["artifact_uri"])
        verdict = holoqa_update_case_verdict(run_id, token, test_id, "PASS", "Fixture state is Complete", artifacts=artifacts)
        assert verdict["completed_cases"] == index

    completed = holoqa_complete_run(run_id, token, "PASS", "All 10 fixture tasks passed")
    assert completed["status"] == "PASS"
    assert len(completed["results"]) == 10
    assert {item["kind"] for item in completed["evidence"]} >= {"screenshot", "dom", "console", "network", "trace", "report_json", "report_html"}
    assert len([item for item in completed["evidence"] if item["kind"] == "screenshot"]) == 10
    assert len([item for item in completed["evidence"] if item["kind"] == "network"]) == 10
    assert all(artifact_exists(uri) for uri in evidence_uris)
    assert all(artifact_exists(item["artifact_path"]) for item in completed["evidence"] if item["artifact_path"].startswith("s3://"))

    api_run = holoqa_get_run_status(run_id)
    assert api_run["completed_cases"] == 10
    assert api_run["total_cases"] == 10
    assert len(api_run["results"]) == 10
    assert any(item["kind"] == "report_json" for item in api_run["evidence"])
    response = TestClient(create_app(database)).get(f"/api/runs/{run_id}")
    assert response.status_code == 200
    assert len(response.json()["results"]) == 10
    print(f"E2E_RUN_ID={run_id}")
